"""Document extraction, text processing, and chunking endpoints for AI Engine."""

import base64
import hashlib
import io
import logging
import math
from typing import Any, Dict, List, Optional

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from pydantic import BaseModel
import pymupdf  # PyMuPDF
import docx
import openpyxl
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.config import settings

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/v1/extract", tags=["Document Extraction"])


class DocumentExtractRequest(BaseModel):
    file_bytes_base64: Optional[str] = None
    filename: str = "document"
    mime_type: str = "text/plain"
    chunk_size: int = 1500
    chunk_overlap: int = 200


class ChunkResult(BaseModel):
    chunk_index: int
    content: str
    token_count: int
    embedding: Optional[List[float]] = None


class DocumentExtractResponse(BaseModel):
    filename: str
    mime_type: str
    page_count: int
    character_count: int
    summary_preview: str
    total_chunks: int
    chunks: List[ChunkResult]


def _generate_fallback_embedding(text: str, dim: int = 1536) -> List[float]:
    """Generates a deterministic, normalized 1536-dimensional embedding vector from text.
    
    Used when no external embedding API key is configured, ensuring pgvector storage
    always receives a valid vector format.
    """
    raw_hash = hashlib.sha512(text.encode("utf-8", errors="ignore")).digest()
    vec = [0.0] * dim
    for i in range(dim):
        byte_val = raw_hash[i % len(raw_hash)]
        # Mix with position to spread values
        pos_factor = math.sin((i + 1) * 0.1)
        vec[i] = (float(byte_val) / 255.0 - 0.5) * pos_factor
    # Normalize vector to unit length (L2 norm)
    norm = math.sqrt(sum(x * x for x in vec)) or 1.0
    return [round(x / norm, 6) for x in vec]


def _compute_embeddings(chunks: List[str]) -> List[List[float]]:
    """Generates 1536-dim vector embeddings for each chunk.
    
    Attempts OpenAI text-embedding-3-small if OPENAI_API_KEY is available,
    otherwise falls back to deterministic normalized embeddings.
    """
    openai_key = settings.__dict__.get("openai_api_key") or None
    import os
    effective_openai_key = openai_key or os.getenv("OPENAI_API_KEY")

    if effective_openai_key:
        try:
            from openai import OpenAI
            client = OpenAI(api_key=effective_openai_key)
            response = client.embeddings.create(
                input=chunks,
                model="text-embedding-3-small"
            )
            return [data.embedding for data in response.data]
        except Exception as exc:
            logger.warning("OpenAI embedding generation failed, using fallback: %s", exc)

    return [_generate_fallback_embedding(chunk) for chunk in chunks]


def extract_text_from_bytes(file_bytes: bytes, filename: str, mime_type: str) -> tuple[str, int]:
    """Extracts raw text and page count from file bytes according to MIME type."""
    full_text = ""
    page_count = 1

    lower_name = filename.lower()
    lower_mime = mime_type.lower()

    try:
        # PDF documents
        if "pdf" in lower_mime or lower_name.endswith(".pdf"):
            doc = pymupdf.open(stream=file_bytes, filetype="pdf")
            page_count = len(doc)
            pages_text = []
            for i, page in enumerate(doc):
                page_str = page.get_text()
                if page_str.strip():
                    pages_text.append(f"[Page {i + 1}]\n{page_str.strip()}")
            full_text = "\n\n".join(pages_text)
            doc.close()

        # Word documents (DOCX)
        elif "wordprocessingml" in lower_mime or lower_name.endswith(".docx"):
            bio = io.BytesIO(file_bytes)
            doc = docx.Document(bio)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            full_text = "\n\n".join(paragraphs)
            page_count = max(1, len(paragraphs) // 10)

        # Excel spreadsheets (XLSX)
        elif "spreadsheetml" in lower_mime or lower_name.endswith(".xlsx"):
            bio = io.BytesIO(file_bytes)
            wb = openpyxl.load_workbook(bio, data_only=True)
            sheet_texts = []
            page_count = len(wb.sheetnames)
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                rows_text = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(cell) for cell in row if cell is not None and str(cell).strip()]
                    if row_vals:
                        rows_text.append(", ".join(row_vals))
                if rows_text:
                    sheet_texts.append(f"--- Sheet: {sheetname} ---\n" + "\n".join(rows_text))
            full_text = "\n\n".join(sheet_texts)

        # Plain text, CSV, Markdown, JSON
        else:
            try:
                full_text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                full_text = file_bytes.decode("latin-1", errors="replace")

    except Exception as exc:
        logger.error("Failed to extract text from %s (%s): %s", filename, mime_type, exc)
        raise HTTPException(status_code=422, detail=f"Text extraction failed: {str(exc)}")

    return full_text.strip(), page_count


def process_and_chunk(
    full_text: str,
    filename: str,
    mime_type: str,
    page_count: int,
    chunk_size: int = 1500,
    chunk_overlap: int = 200,
) -> DocumentExtractResponse:
    """Chunks text and returns structured response."""
    if not full_text:
        return DocumentExtractResponse(
            filename=filename,
            mime_type=mime_type,
            page_count=page_count,
            character_count=0,
            summary_preview="",
            total_chunks=0,
            chunks=[],
        )

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )

    raw_chunks = splitter.split_text(full_text)
    if not raw_chunks and full_text:
        raw_chunks = [full_text]

    embeddings = _compute_embeddings(raw_chunks)

    chunk_results: List[ChunkResult] = []
    for i, chunk_text in enumerate(raw_chunks):
        approx_tokens = max(1, len(chunk_text) // 4)
        emb = embeddings[i] if i < len(embeddings) else None
        chunk_results.append(
            ChunkResult(
                chunk_index=i,
                content=chunk_text,
                token_count=approx_tokens,
                embedding=emb,
            )
        )

    summary_preview = full_text[:400].replace("\n", " ").strip()
    if len(full_text) > 400:
        summary_preview += "..."

    return DocumentExtractResponse(
        filename=filename,
        mime_type=mime_type,
        page_count=page_count,
        character_count=len(full_text),
        summary_preview=summary_preview,
        total_chunks=len(chunk_results),
        chunks=chunk_results,
    )


@router.post("/document", response_model=DocumentExtractResponse)
async def extract_document_from_json(payload: DocumentExtractRequest) -> DocumentExtractResponse:
    """Extracts text, splits into chunks, and computes embeddings from base64 document bytes."""
    if not payload.file_bytes_base64:
        raise HTTPException(status_code=400, detail="file_bytes_base64 is required")

    try:
        file_bytes = base64.b64decode(payload.file_bytes_base64)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid base64 encoding: {str(exc)}")

    full_text, page_count = extract_text_from_bytes(file_bytes, payload.filename, payload.mime_type)
    return process_and_chunk(
        full_text=full_text,
        filename=payload.filename,
        mime_type=payload.mime_type,
        page_count=page_count,
        chunk_size=payload.chunk_size,
        chunk_overlap=payload.chunk_overlap,
    )


@router.post("/upload", response_model=DocumentExtractResponse)
async def extract_document_from_upload(
    file: UploadFile = File(...),
    chunk_size: int = Form(1500),
    chunk_overlap: int = Form(200),
) -> DocumentExtractResponse:
    """Extracts text, splits into chunks, and computes embeddings from multipart uploaded file."""
    file_bytes = await file.read()
    filename = file.filename or "uploaded_file"
    mime_type = file.content_type or "application/octet-stream"

    full_text, page_count = extract_text_from_bytes(file_bytes, filename, mime_type)
    return process_and_chunk(
        full_text=full_text,
        filename=filename,
        mime_type=mime_type,
        page_count=page_count,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
    )
